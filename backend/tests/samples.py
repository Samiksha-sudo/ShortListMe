"""Builders for sample CV DOCX files used by the parser tests."""
import io

from docx import Document


def _save(doc: Document) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _bullet(doc_or_cell, text: str):
    return doc_or_cell.add_paragraph(text, style="List Bullet")


def single_column() -> bytes:
    d = Document()
    d.add_paragraph("Jane Doe")
    d.add_paragraph("jane@example.com | +44 7700 900123 | London, UK")
    d.add_paragraph("Backend Software Engineer")
    d.add_paragraph("Summary")
    d.add_paragraph("Backend engineer with 4 years of experience building web services.")
    d.add_paragraph("Experience")
    d.add_paragraph("Software Engineer II, FoodCo | 2022 - Present")
    _bullet(d, "Built REST APIs for the order service handling 2M requests/day.")
    _bullet(d, "Worked on the deployment pipeline to speed up releases.")
    d.add_paragraph("Software Engineer, ShopLab | 2020 - 2022")
    _bullet(d, "Developed React dashboards used by 300 support agents.")
    d.add_paragraph("Skills")
    d.add_paragraph("Node.js, TypeScript, React, PostgreSQL, Docker, AWS")
    d.add_paragraph("Education")
    d.add_paragraph("BSc Computer Science, University of Leeds, 2016 - 2019")
    return _save(d)


def glyph_bullets() -> bytes:
    """Bullets typed as text glyphs instead of Word list formatting."""
    d = Document()
    d.add_paragraph("John Smith")
    d.add_paragraph("john@example.com")
    d.add_paragraph("Senior Data Engineer")
    d.add_paragraph("WORK EXPERIENCE")
    d.add_paragraph("Data Engineer, DataCo | 2019 - Present")
    d.add_paragraph("• Built ETL pipelines processing 5TB per day.")
    d.add_paragraph("- Reduced warehouse costs by 30%.")
    return _save(d)


def table_based() -> bytes:
    """Two-column layout done with a table: sidebar (contact, skills) + main (experience)."""
    d = Document()
    t = d.add_table(rows=1, cols=2)
    side, main = t.rows[0].cells
    side.paragraphs[0].text = "Jane Doe"
    side.add_paragraph("jane@example.com")
    side.add_paragraph("Skills")
    side.add_paragraph("Python, FastAPI, SQL")
    main.paragraphs[0].text = "Experience"
    main.add_paragraph("Backend Developer, Acme Ltd | 2021 - Present")
    _bullet(main, "Designed internal APIs used by 12 teams.")
    return _save(d)


def with_header_footer() -> bytes:
    d = Document()
    d.sections[0].header.paragraphs[0].text = "Jane Doe - CV"
    d.sections[0].footer.paragraphs[0].text = "Page 1"
    d.add_paragraph("Jane Doe")
    d.add_paragraph("Software Engineer")
    return _save(d)
