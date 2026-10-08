import io

from docx import Document

from tests import samples


def _upload(client, data):
    r = client.post("/v1/cvs", files={"file": ("cv.docx", data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert r.status_code == 201
    return r.json()["cv_id"]


def _blocks(client, cv_id):
    return client.get(f"/v1/cvs/{cv_id}").json()["blocks"]


def test_export_applies_edit_and_keeps_rest(client):
    cv_id = _upload(client, samples.single_column())
    bullet = next(b for b in _blocks(client, cv_id) if b["type"] == "bullet")
    r = client.post(f"/v1/cvs/{cv_id}/export.docx", json={"edits": {bullet["id"]: "Built REST APIs for the order service."}})
    assert r.status_code == 200
    assert "attachment" in r.headers["content-disposition"]
    texts = [p.text for p in Document(io.BytesIO(r.content)).paragraphs]
    assert "Built REST APIs for the order service." in texts
    assert "Jane Doe" in texts
    assert not any("2M requests/day" in t for t in texts)


def test_export_without_edits_returns_original_text(client):
    cv_id = _upload(client, samples.single_column())
    r = client.post(f"/v1/cvs/{cv_id}/export.docx", json={"edits": {}})
    assert r.status_code == 200
    assert "Jane Doe" in [p.text for p in Document(io.BytesIO(r.content)).paragraphs]


def test_export_restores_glyph_prefix(client):
    cv_id = _upload(client, samples.glyph_bullets())
    bullet = next(b for b in _blocks(client, cv_id) if b["type"] == "bullet")
    r = client.post(f"/v1/cvs/{cv_id}/export.docx", json={"edits": {bullet["id"]: "Built ETL pipelines."}})
    texts = [p.text for p in Document(io.BytesIO(r.content)).paragraphs]
    assert f"{bullet['locator']['prefix']}Built ETL pipelines." in texts


def test_export_rejects_locked_and_unknown_blocks(client):
    cv_id = _upload(client, samples.single_column())
    name = next(b for b in _blocks(client, cv_id) if b["type"] == "name")
    assert client.post(f"/v1/cvs/{cv_id}/export.docx", json={"edits": {name["id"]: "X"}}).status_code == 422
    assert client.post(f"/v1/cvs/{cv_id}/export.docx", json={"edits": {"b_999": "X"}}).status_code == 422


def test_export_unknown_cv(client):
    assert client.post("/v1/cvs/nope/export.docx", json={"edits": {}}).status_code == 404


def test_export_pdf(client):
    import pytest
    pytest.importorskip("reportlab")
    cv_id = _upload(client, samples.single_column())
    bullet = next(b for b in _blocks(client, cv_id) if b["type"] == "bullet")
    r = client.post(f"/v1/cvs/{cv_id}/export.pdf", json={"edits": {bullet["id"]: "Built REST APIs."}})
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF")
    assert client.post(f"/v1/cvs/{cv_id}/export.pdf", json={"edits": {"b_999": "x"}}).status_code == 422
