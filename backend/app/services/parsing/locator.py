from docx.document import Document
from docx.text.paragraph import Paragraph

from app.schemas.cv import Locator


def locate(doc: Document, loc: Locator) -> Paragraph:
    """Find the paragraph a block points to (used by the exporter)."""
    if loc.table:
        t, r, c = loc.table
        return doc.tables[t].rows[r].cells[c].paragraphs[loc.p]
    if loc.part == "body":
        return doc.paragraphs[loc.p]
    kind, idx = loc.part.split(":")
    section = doc.sections[int(idx)]
    container = section.header if kind == "header" else section.footer
    return container.paragraphs[loc.p]
